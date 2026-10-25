"""Generate the Excel import template and the filled example from excel_import/levee_excel.py.

    python tools/gen_import_template.py      # needs openpyxl (maintainers only; users never run this)

Writes excel_import/levee_import_template.xlsx (empty) and excel_import/levee_import_example.xlsx
(the 24 synthetic sections of the simplified batch page, 3 rows with deliberate mistakes).
example_requests() and example_rows() need no third-party package; the tests use them.
"""

from __future__ import annotations

import copy
import json
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "excel_import"))
sys.path.insert(0, str(ROOT))

import levee_excel as L  # noqa: E402

OUT_DIR = ROOT / "excel_import"
TEMPLATE = OUT_DIR / "levee_import_template.xlsx"
EXAMPLE = OUT_DIR / "levee_import_example.xlsx"
BAD_ROWS = {3: "crest_width_zero", 8: "fill_quality_typo", 17: "no_foundation_gradient"}   # positions in the batch
FILL = {"basic": "D9E1F2", "auto": "DDEBF7", "fill": "FFF2CC", "optional": "EDEDED", "detail": "D9E1F2"}
A6_AS_RECORDS = {100: [], 80: [("蚁穴", "已处理")], 60: [("裂缝", "未处理")], 0: [("蚁穴", "未处理")]}
A9_NAMES = {100: "无影响", 70: "轻微影响", 0: "存在危险"}
B2_NAMES = {100: "符合", 60: "待复核", 0: "不符合"}


def example_requests():
    """The batch of the simplified page: 21 synthetic sections with data-supported inputs + 3 bad copies."""
    reqs = json.loads((ROOT / "tests" / "fixtures" / "synthetic_875_requests.json").read_text(encoding="utf-8"))
    rng = random.Random(7)
    picked = [copy.deepcopy(reqs[i]) for i in sorted(rng.sample(range(875), 21))]
    rng.shuffle(picked)
    items = []
    for k, q in enumerate(picked):
        for c in ("A4", "B3", "C2", "D5", "E1", "E2", "E3", "E4", "E5"):
            q["qualitative"].pop(c, None)
        q["survey"].pop("soft_foundation_settlement_cm")
        q["survey"].pop("soft_foundation_settlement_diff_cm")
        q["survey"]["soft_foundation_status"] = "unknown"
        q.pop("inspection")
        if k % 3 == 0:
            q["computation"].pop("back_slope_fos")
            q["computation"].pop("front_slope_fos")
        if k % 4 == 1:
            q["survey"].pop("scour_depth_m")
        items.append(q)
    bad = {3: copy.deepcopy(items[0]), 8: copy.deepcopy(items[1]), 17: copy.deepcopy(items[2])}
    bad[3]["section"]["crest_width_m"] = 0
    bad[17]["computation"].pop("foundation_gradient")
    for pos in sorted(bad):
        items.insert(pos, bad[pos])
    for k, q in enumerate(items):
        q["request_id"] = f"batch-{k + 1:04d}"
        q["section"]["section_id"] = f"K{(k * 200) // 1000}+{(k * 200) % 1000:03d}"
    return items


def example_rows():
    """Template cells for example_requests(): (main rows, danger rows, event rows, other-indicator rows)."""
    main, dangers, other = [], [], []
    for k, q in enumerate(example_requests()):
        s, comp, survey, rec, qual = q["section"], q["computation"], q["survey"], q["records"], q["qualitative"]
        sid = s["section_id"]
        a6 = qual.get("A6", 100)
        row = {
            "section_id": sid, "levee_type": "湖堤" if s["is_lake_levee"] else "河堤",
            "crest_elevation": s["crest_elevation_m"], "layer_name": "粉细砂（合成示例）",
            "layer_gs": s["control_layer"]["Gs"], "layer_e": s["control_layer"]["e"],
            "grad_foundation": comp.get("foundation_gradient"), "grad_body": comp["body_gradient"],
            "grad_contact": comp["contact_gradient"], "design_flood_level": s["design_flood_level_m"],
            "crest_width": s["crest_width_m"], "outer_slope": f"1:{s['outer_slope_ratio']:g}",
            "inner_slope": f"1:{s['inner_slope_ratio']:g}", "foundation_class": s["foundation_class"],
            "fill_quality": "好" if BAD_ROWS.get(k) == "fill_quality_typo" else rec["fill_quality"].replace("质量", ""),
            "flood_standard": B2_NAMES[qual["B2"]],
            "structure": "无" if rec["crossing_structure_class"] == "不涉及" else rec["crossing_structure_class"],
            "over_crossing": A9_NAMES[qual["A9"]], "dangers": "有" if A6_AS_RECORDS[a6] else "无",
            "back_fos": comp.get("back_slope_fos"), "front_fos": comp.get("front_slope_fos"),
            "scour_depth": survey.get("scour_depth_m"),
        }
        main.append(row)
        for typ, done in A6_AS_RECORDS[a6]:
            dangers.append({"section_id": sid, "type": typ, "reinforced": done})
        if rec["crossing_structure_class"] != "不涉及":
            for code in L.STRUCTURE_INDICATORS:
                other.append({"section_id": sid, "code": code, "value": qual[code]})
    return main, dangers, [], other


def _write(path, rows=None):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")
    main_rows, danger_rows, event_rows, other_rows = rows or ([], [], [], [])
    wb = Workbook()
    help_ws = wb.active
    help_ws.title = L.SHEET_HELP

    def header(ws, row, col, c, fill):
        cell = ws.cell(row=row, column=col, value=c.header + ("*" if c.required else ""))
        cell.font = Font(bold=True, color="C00000" if c.required else "000000")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment, cell.border = wrap, border
        cell.comment = Comment(L.fill_hint(c), "导入模板", width=260, height=110)
        ws.column_dimensions[get_column_letter(col)].width = max(10, min(18, len(c.header) * 2 + 4))

    def validations(ws, columns, first, last):
        for col, c in enumerate(columns, 1):
            letter = get_column_letter(col)
            common = dict(allow_blank=True, showInputMessage=True, prompt=L.fill_hint(c)[:250], promptTitle=c.header[:30])
            if c.kind in ("choice", "structure", "dangers", "soft", "inspection", "bool"):
                shown = list(_BOOL_SHOWN) if c.kind == "bool" else [k for k in c.options if not k[0].isdigit()]
                dv = DataValidation(type="list", formula1='"' + ",".join(shown) + '"', showErrorMessage=True,
                                    errorStyle="stop", error="请从下拉列表中选择", errorTitle=c.header, **common)
            elif c.kind in ("number", "integer"):
                low = "-1000" if c.key in ("crest_elevation", "design_flood_level") else "0"
                dv = DataValidation(type="whole" if c.kind == "integer" else "decimal", operator="between",
                                    formula1=low, formula2="100000", showErrorMessage=True, errorStyle="warning",
                                    error="这里应填数字；没有资料请留空", errorTitle=c.header, **common)
            else:
                dv = DataValidation(**common)
            dv.add(f"{letter}{first}:{letter}{last}")
            ws.add_data_validation(dv)

    # ---- 断面: a band row naming the column groups, then the header row
    ws = wb.create_sheet(L.SHEET_MAIN)
    col = 1
    for key, title in L.GROUPS:
        cols = [c for c in L.COLUMNS if c.group == key]
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(cols) - 1)
        band = ws.cell(row=1, column=col, value=title)
        band.font, band.alignment = Font(bold=True), wrap
        for j in range(col, col + len(cols)):
            ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor=FILL[key])
            ws.cell(row=1, column=j).border = border
        for c in cols:
            header(ws, 2, col, c, FILL[key])
            col += 1
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 36
    ws.freeze_panes = "B3"
    validations(ws, L.COLUMNS, 3, 1002)
    for j, c in enumerate(L.COLUMNS, 1):
        if c.kind in ("ratio", "text"):    # text format, so that Excel keeps “1:3” instead of turning it into a time
            ws.column_dimensions[get_column_letter(j)].number_format = "@"
    for r, values in enumerate(main_rows, 3):
        for j, c in enumerate(L.COLUMNS, 1):
            if values.get(c.key) is not None:
                cell = ws.cell(row=r, column=j, value=values[c.key])
                if c.kind in ("ratio", "text"):
                    cell.number_format = "@"

    # ---- detail sheets
    for name, columns, data in ((L.SHEET_DANGERS, L.DANGER_COLUMNS, danger_rows),
                                (L.SHEET_EVENTS, L.EVENT_COLUMNS, event_rows),
                                (L.SHEET_OTHER, L.OTHER_COLUMNS, other_rows)):
        ds = wb.create_sheet(name)
        for j, c in enumerate(columns, 1):
            header(ds, 1, j, c, FILL["detail"])
        ds.row_dimensions[1].height = 30
        ds.freeze_panes = "A2"
        validations(ds, columns, 2, 1001)
        for r, values in enumerate(data, 2):
            for j, c in enumerate(columns, 1):
                v = values.get(c.key)
                if isinstance(v, bool):
                    v = ("已处理" if v else "未处理") if c.key == "reinforced" else ("是" if v else "否")
                if v is not None:
                    ds.cell(row=r, column=j, value=v)

    # ---- 填写说明
    lines = [
        ("堤防安全评价 · 断面资料导入模板", "title"),
        ("", None),
        ("怎么填", "h"),
        ("1. 在“断面”表一行填一个断面（评价单元）。表头带 * 的列必填；选填列没有资料就留空，不要用假设值补齐。", None),
        ("2. 浅蓝色的“前两步带入”列是地质建模和渗流计算的结果，平台可按桩号自动填入；单独使用本模板时手工填写。", None),
        ("3. 黄色的“需要填写”列一般可从设计报告、地勘报告和工程台账查到；下拉选项之外的内容会被拒绝。", None),
        ("4. 有历史险情、巡检确认的险情或其他指标结论时，分别填在“历史险情”“巡检事件”“其他指标”表，用桩号对应。", None),
        ("5. 导入后，有问题的行会提示行号和改法，其余行照常评价；没填的选填项对应的限级检查会显示为“未做”。", None),
        ("6. 评价等级 A/B/C 是本评价模型的研究分级，不等同于规范中的安全类别。", None),
        ("", None),
        ("各列说明", "h"),
    ]
    for r, (text, kind) in enumerate(lines, 1):
        cell = help_ws.cell(row=r, column=1, value=text)
        if kind == "title":
            cell.font = Font(bold=True, size=15)
        elif kind == "h":
            cell.font = Font(bold=True, size=12)
    r = len(lines) + 1
    for j, h in enumerate(("表", "分组", "列", "必填", "怎么填", "对应接口字段"), 1):
        cell = help_ws.cell(row=r, column=j, value=h)
        cell.font, cell.fill, cell.border = Font(bold=True), PatternFill("solid", fgColor="D9D9D9"), border
    groups = dict(L.GROUPS)
    table = [(L.SHEET_MAIN, groups[c.group], c, FILL[c.group]) for c in L.COLUMNS]
    for name, columns in L.DETAIL_SHEETS.items():
        table += [(name, "", c, FILL["detail"]) for c in columns]
    for name, group, c, fill in table:
        r += 1
        api_path = c.path or ""
        if c.kind == "structure":
            api_path += "；选“无”时 qualitative.A8/B4/C6/D4 = 不涉及"
        values = (name, group, c.header, "是" if c.required else "", L.fill_hint(c), api_path)
        for j, v in enumerate(values, 1):
            cell = help_ws.cell(row=r, column=j, value=v)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 3:
                cell.fill = PatternFill("solid", fgColor=fill)
    for letter, width in zip("ABCDEF", (10, 20, 16, 6, 60, 36)):
        help_ws.column_dimensions[letter].width = width
    wb.active = 0
    wb.save(path)


_BOOL_SHOWN = ("是", "否")


def main():
    _write(TEMPLATE)
    _write(EXAMPLE, example_rows())
    print(f"wrote {TEMPLATE.relative_to(ROOT)}, {EXAMPLE.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
