"""Inventory and provenance audit for the newly delivered inspection package."""

from __future__ import annotations

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image, ExifTags

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "巡检识别"
OUT = ROOT / "分析工作区" / "巡检识别审计"


def sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader(); w.writerows(rows)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    files = [p for p in SRC.rglob("*") if p.is_file()]
    manifest = []
    for p in sorted(files):
        rel = str(p.relative_to(ROOT))
        suffix = p.suffix.lower()
        if suffix in {".jpg", ".jpeg", ".png"}:
            role = "image"
        elif suffix == ".txt": role = "detection_label"
        elif suffix == ".xlsx": role = "review_table"
        elif suffix == ".pdf": role = "radiometric_report"
        elif suffix == ".docx": role = "package_description"
        else: role = "other"
        manifest.append({"relative_path": rel, "bytes": p.stat().st_size, "sha256": sha(p), "role": role})
    write_csv(OUT / "巡检资料文件清单.csv", manifest, ["relative_path", "bytes", "sha256", "role"])

    # Extract narrative package description and radiometric report text as UTF-8 analysis copies.
    doc = Document(SRC / "巡检识别资料汇总说明.docx")
    doc_text = "\n".join(x.text for x in doc.paragraphs)
    (OUT / "资料汇总说明_提取文本.txt").write_text(doc_text, encoding="utf-8")
    pdf = PdfReader(SRC / "第二批补充资料" / "01_原始辐射温度" / "原始图像辐射温度报告.pdf")
    pdf_text = "\n\n".join(page.extract_text() or "" for page in pdf.pages)
    (OUT / "原始图像辐射温度报告_提取文本.txt").write_text(pdf_text, encoding="utf-8")

    # Review table: retain nonempty records and all original headers.
    xlsx = SRC / "第一批优先资料" / "04_可疑点复核记录" / "可疑点复核记录表.xlsx"
    wb = load_workbook(xlsx, data_only=False)
    table_rows = []
    for ws in wb.worksheets:
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1):
            vals = [clean(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            if any(v not in (None, "") and not (isinstance(v, str) and v.startswith("=_xlfn.DISPIMG")) for v in vals):
                table_rows.append({"sheet": ws.title, **{str(headers[i] or f"col_{i+1}"): vals[i] for i in range(len(headers))}})
    write_csv(OUT / "可疑点复核记录_非空行.csv", table_rows, list(table_rows[0]))

    # Detection label files: YOLO-style class x y w h confidence (preserve raw line).
    label_rows = []
    for p in sorted(SRC.rglob("*.txt")):
        group = "人工管涌验证" if "人造管涌" in str(p) else "匿名堤防现场巡检"
        raw = p.read_text(encoding="utf-8-sig", errors="replace").strip()
        lines = [line for line in raw.splitlines() if line.strip()]
        for line_no, line in enumerate(lines, 1):
            fields = line.split()
            label_rows.append({"group": group, "file": str(p.relative_to(ROOT)), "line": line_no, "n_fields": len(fields), "raw": line,
                               "class": fields[0] if fields else "", "x": fields[1] if len(fields) > 1 else "", "y": fields[2] if len(fields) > 2 else "", "w": fields[3] if len(fields) > 3 else "", "h": fields[4] if len(fields) > 4 else "", "confidence": fields[5] if len(fields) > 5 else ""})
    write_csv(OUT / "检测标签逐条记录.csv", label_rows, list(label_rows[0]))

    # Cross-check the nine field review records against TXT confidence values.
    label_by_stem = defaultdict(list)
    for row in label_rows:
        label_by_stem[Path(row["file"]).stem].append(row)
    cross = []
    for row in table_rows:
        filename = str(row.get("文件名") or "")
        stem = Path(filename).stem
        labels = label_by_stem.get(stem, [])
        txt_conf = [float(x["confidence"]) for x in labels if x["confidence"]]
        table_conf = float(row["置信度"]) if row.get("置信度") not in (None, "") else None
        cross.append({"filename": filename, "review_result": row.get("复核结果"), "table_confidence": table_conf,
                      "label_confidence": txt_conf, "confidence_delta": round(table_conf - txt_conf[0], 6) if txt_conf and table_conf is not None else None,
                      "confidence_match": bool(txt_conf and table_conf is not None and abs(table_conf - txt_conf[0]) < 0.005),
                      "gps": row.get("经纬度坐标（WGS84）"), "capture_time": row.get("拍摄时间")})
    write_csv(OUT / "现场复核与检测输出交叉核对.csv", cross, list(cross[0]))

    # Image dimensions and EXIF availability; use one row per image, then summarize.
    image_rows = []
    for p in files:
        if p.suffix.lower() not in {".jpg", ".jpeg", ".png"}:
            continue
        try:
            with Image.open(p) as im:
                exif = im.getexif()
                image_rows.append({"relative_path": str(p.relative_to(ROOT)), "width": im.width, "height": im.height,
                                   "mode": im.mode, "format": im.format, "exif_tags": len(exif),
                                   "has_gps_tag": 34853 in exif, "has_datetime_tag": 306 in exif})
        except Exception as exc:
            image_rows.append({"relative_path": str(p.relative_to(ROOT)), "width": "", "height": "", "mode": "ERROR", "format": "", "exif_tags": "", "has_gps_tag": "", "has_datetime_tag": "", "error": str(exc)})
    write_csv(OUT / "图像尺寸与EXIF清单.csv", image_rows, list(image_rows[0]))

    def count_images(part: str):
        return len([p for p in files if part in str(p) and p.suffix.lower() in {".jpg", ".jpeg", ".png"}])
    summary = {
        "file_counts": dict(Counter(p.suffix.lower() for p in files)),
        "image_counts": {"field_raw_and_outputs": count_images("01_原始图像及视频"), "artificial_validation": count_images("人造管涌"), "thermal_raw": count_images("原始辐射温度图像"), "track_screenshots": count_images("飞行轨迹与时间戳")},
        "label_files": len({r["file"] for r in label_rows}), "label_rows": len(label_rows),
        "label_groups": dict(Counter(r["group"] for r in label_rows)),
        "label_fields": dict(Counter(r["n_fields"] for r in label_rows)),
        "label_classes": dict(Counter(r["class"] for r in label_rows)),
        "review_nonempty_rows": len(table_rows), "review_columns": list(table_rows[0]) if table_rows else [],
        "review_records_with_algorithm_label": sum(bool(r.get("识别类型")) for r in table_rows),
        "review_records_confirmed": dict(Counter(str(r.get("复核结果")) for r in table_rows)),
        "review_confidence_crosscheck": {"records": len(cross), "matches": sum(r["confidence_match"] for r in cross), "mismatches": sum(not r["confidence_match"] for r in cross)},
        "image_dimensions": {f"{r[0]}x{r[1]}": n for r, n in Counter((r["width"], r["height"]) for r in image_rows).items()},
        "images_with_gps_exif": sum(bool(r["has_gps_tag"]) for r in image_rows),
        "images_with_datetime_exif": sum(bool(r["has_datetime_tag"]) for r in image_rows),
        "pdf_pages": len(pdf.pages),
        "pdf_text_chars": len(pdf_text),
        "description_text_chars": len(doc_text),
        "status": "inventory and structure audit; no image-level accuracy claim yet",
    }
    (OUT / "巡检识别审计摘要.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
